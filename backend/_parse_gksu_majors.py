# -*- coding: utf-8 -*-
"""Robustly parse GKS-U University Information xlsx: detect header columns dynamically."""
import glob
import os
import json
import openpyxl

BASE = r"C:\Users\USER\camnemi-crm\backend\_gksu_univinfo"
OUT = r"C:\Users\USER\camnemi-crm\backend\_gksu_uic_majors.json"

UIC = ["Ajou University", "Daegu University", "Dong-A University", "Inje University",
       "Keimyung University", "Konyang University", "Kookmin University",
       "Korea University of Technology and Education", "Sungshin Women",
       "Yeungnam University"]
ASSOC = ["Dong-eui Institute of Technology", "Dong Eui Institute of Technology",
         "Gumi University", "Hanyang Women", "Hosan University",
         "Korea University of Media Arts", "Kyungbok University", "Osan University",
         "Vision College of Jeonju", "Yeungjin University"]


def find_folder(name):
    for d in glob.glob(os.path.join(BASE, "*")):
        if os.path.isdir(d) and name.lower() in os.path.basename(d).lower():
            return d
    return None


def find_header(ws):
    """Return {colname: idx} if this sheet has the departments table header, else None."""
    for row in ws.iter_rows(values_only=True):
        cells = [str(c).strip() if c else "" for c in row]
        joined = "|".join(cells)
        if "모집단위" in joined and "Department" in joined and "TOPIK" in joined:
            idx = {}
            for i, c in enumerate(cells):
                cl = c.lower()
                if "모집단위" in c:
                    idx["kr_dept"] = i
                elif "department" in cl:
                    idx["en_dept"] = i
                elif "학과계열" in c:
                    idx["field_kr"] = i
                elif "field of study" in cl:
                    idx["field_en"] = i
                elif "medium" in cl:
                    idx["medium"] = i
                elif "required topik" in cl:
                    idx["topik"] = i
                elif "univ. track" in cl:
                    idx["track"] = i
                elif "embassy track" in cl:
                    idx["embassy"] = i
                elif "campus" in cl:
                    idx["campus"] = i
            if "kr_dept" in idx and "en_dept" in idx:
                return idx
    return None


def main():
    results = {}
    for school in UIC + ASSOC:
        folder = find_folder(school)
        if not folder:
            continue
        xlsx = glob.glob(os.path.join(folder, "*.xlsx"))
        if not xlsx:
            continue
        wb = openpyxl.load_workbook(xlsx[0], read_only=True, data_only=True)
        entries = []
        for ws in wb.worksheets:
            idx = find_header(ws)
            if not idx:
                continue
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() if c is not None else "" for c in row]
                joined = " ".join(cells)
                if "Available Departments" in joined or joined.startswith("No."):
                    continue
                kr = cells[idx["kr_dept"]] if len(cells) > idx["kr_dept"] else ""
                en = cells[idx["en_dept"]] if len(cells) > idx["en_dept"] else ""
                if not kr and not en:
                    continue
                entries.append({
                    "kr": kr, "en": en,
                    "field_kr": cells[idx.get("field_kr", 0)] if idx.get("field_kr") is not None and len(cells) > idx.get("field_kr", 99) else "",
                    "field_en": cells[idx.get("field_en", 0)] if idx.get("field_en") is not None and len(cells) > idx.get("field_en", 99) else "",
                    "medium": cells[idx.get("medium", 0)] if idx.get("medium") is not None and len(cells) > idx.get("medium", 99) else "",
                    "topik": cells[idx.get("topik", 0)] if idx.get("topik") is not None and len(cells) > idx.get("topik", 99) else "",
                    "track": cells[idx.get("track", 0)] if idx.get("track") is not None and len(cells) > idx.get("track", 99) else "",
                    "embassy": cells[idx.get("embassy", 0)] if idx.get("embassy") is not None and len(cells) > idx.get("embassy", 99) else "",
                    "campus": cells[idx.get("campus", 0)] if idx.get("campus") is not None and len(cells) > idx.get("campus", 99) else "",
                })
        results[school] = entries
        print(f"{school}: {len(entries)} depts")
    json.dump(results, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("saved:", OUT)


if __name__ == "__main__":
    main()
